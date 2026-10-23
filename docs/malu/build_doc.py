from __future__ import annotations

import argparse
import math
import re
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
NAVY = '163550'
TEAL = '087F8C'
PALE = 'EDF5F8'
GRAY = '526473'
INK = '213342'
FONT = 'Noto Sans CJK SC'


def rfont(run, size=None, bold=None, color=None, mono=False):
    run.font.name = 'DejaVu Sans Mono' if mono else FONT
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'DejaVu Sans Mono' if mono and key != 'eastAsia' else FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    prop = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), fill)
    prop.append(el)


def prevent_row_split(row):
    el = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(el)


def repeat_header(row):
    el = OxmlElement('w:tblHeader')
    row._tr.get_or_add_trPr().append(el)


def field(paragraph, instruction):
    run = paragraph.add_run()
    start = OxmlElement('w:fldChar')
    start.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText')
    code.set(qn('xml:space'), 'preserve')
    code.text = instruction
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in [start, code, end]:
        run._r.append(el)
    rfont(run, size=8.5, color=GRAY)


def hyperlink(paragraph, text, target=None, anchor=None):
    link = OxmlElement('w:hyperlink')
    if anchor:
        link.set(qn('w:anchor'), anchor)
    else:
        from docx.opc.constants import RELATIONSHIP_TYPE
        rel = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), TEAL)
    props.append(color)
    fonts = OxmlElement('w:rFonts')
    for k in ['ascii', 'hAnsi', 'eastAsia']:
        fonts.set(qn('w:' + k), FONT)
    props.append(fonts)
    r.append(props)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def inline(paragraph, text, *, size=None, base_color=None):
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            rfont(run, size=size, bold=True, color=base_color or INK)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            rfont(run, size=size or 9.2, mono=True, color=base_color or NAVY)
        elif (m := re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)):
            hyperlink(paragraph, m.group(1), target=m.group(2))
        else:
            run = paragraph.add_run(token)
            rfont(run, size=size, color=base_color or INK)


def clean_formula(text):
    pairs = {
        r'\mathrm{': '', r'\text{': '', r'\operatorname{': '',
        r'\times': ' × ', r'\cdot': ' · ', r'\sum': 'Σ',
        r'\leq': '≤', r'\geq': '≥', r'\neq': '≠', r'\infty': '∞',
        r'\rightarrow': ' → ', r'\lceil': 'ceil(', r'\rceil': ')',
        r'\lfloor': 'floor(', r'\rfloor': ')', r'\log_2': 'log₂',
        r'\qquad': '   ', r'\quad': '  ', r'\,': ' ',
        r'\left': '', r'\right': '', r'\!': '',
    }
    for a, b in pairs.items():
        text = text.replace(a, b)
    return text.strip()


def add_table(doc, rows):
    cols = max(map(len, rows))
    rows = [row + [''] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    weights = []
    for c in range(cols):
        lengths = sorted(len(re.sub(r'[`*]', '', row[c])) for row in rows)
        typical = lengths[min(len(lengths)-1, int(len(lengths)*0.7))]
        weights.append(math.sqrt(max(5, min(typical, 70))))
    widths = [17.0 * w / sum(weights) for w in weights]
    for c, width in zip(table.columns, widths):
        c.width = Cm(width)
    for i, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if i == 0:
            repeat_header(row)
        for j, text in enumerate(values):
            cell = row.cells[j]
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '75')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tc_pr.append(margins)
            shade(cell, NAVY if i == 0 else ('F0F6F8' if i % 2 else 'FFFFFF'))
            for k, part in enumerate(text.replace('<br>', '\n').split('\n')):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                inline(p, part, size=8.6 if cols >= 6 else 9.0,
                       base_color='FFFFFF' if i == 0 else INK)
                if i == 0:
                    for run in p.runs:
                        run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def bookmark(paragraph, ident):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(ident))
    start.set(qn('w:name'), 'section_' + str(ident))
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(ident))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_figure(doc, path, caption):
    with Image.open(path) as im:
        width, height = im.size
    displayed_width = 17.0
    if displayed_width * height / width > 19.0:
        displayed_width = 19.0 * width / height
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(displayed_width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    inline(p, caption, size=9.2, base_color=GRAY)


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.85)
    section.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = FONT, Pt(10.3)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for level, size in [(1, 15.5), (2, 12.3), (3, 11.1)]:
        style = doc.styles['Heading ' + str(level)]
        style.font.name, style.font.size = FONT, Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
        style.font.bold = True
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    inline(header, 'MALU  /  MATRIX DOT PRODUCT UNIT', size=8.2, base_color=GRAY)
    footer = section.footer.paragraphs[0]
    inline(footer, '设计评审稿  ·  v1.1                                      ', size=8.0, base_color=GRAY)
    field(footer, ' PAGE ')
    inline(footer, ' / ', size=8.0, base_color=GRAY)
    field(footer, ' NUMPAGES ')
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)


def cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    inline(p, 'MATRIX ALU', size=33, base_color=TEAL)
    p = doc.add_paragraph()
    run = p.add_run('MALU 设计文档')
    rfont(run, size=27, bold=True, color=NAVY)
    p = doc.add_paragraph()
    inline(p, '多格式向量点积  ·  SystemVerilog 实现规格', size=13, base_color=GRAY)
    doc.add_paragraph()
    add_table(doc, [
        ['文档属性', '内容'],
        ['版本 / 日期', 'v1.1 设计评审稿 / 2026-10-23（Asia/Shanghai）'],
        ['功能边界', 'D = C + Dot(A, B)，只提供矩阵向量点积事务'],
        ['接口重点', 'A/B/C 数据；独立数据格式与 scale 格式；运行时 BLOCK16 / BLOCK32'],
        ['性能目标', 'Rubin dense 参考配比已核验；逐格式 / scale / BLOCK 映射继续细化'],
        ['当前阶段', '文档与工程骨架发布；未编写 MALU RTL、未进行功能仿真或综合'],
    ])
    p = doc.add_paragraph()
    inline(p, '阅读说明', size=12, base_color=TEAL)
    p = doc.add_paragraph()
    inline(p, '正文明确区分 ISA 要求、文献观察和项目设计选择。PTX 与附件证据已核验的条目给出页码；未公开或尚未取得的内容保留待决状态。', size=10)
    p = doc.add_paragraph()
    inline(p, '重要更正：上传 PTX 原文的 Ti16 类型名为 .s1z4m11；与此前会话记录的 e1z4m11 不同。具体数值语义以正文的原文证据为准。', size=10, base_color=NAVY)
    doc.add_page_break()


def build(source, destination):
    raw = source.read_text(encoding='utf-8')
    lines = raw.splitlines()
    doc = Document()
    configure(doc)
    doc.core_properties.title = 'MALU 多格式向量点积设计文档'
    doc.core_properties.subject = 'PTX ISA 9.4, Matrix ALU, Ti16, block scaling, Rubin throughput target'
    doc.core_properties.author = 'MALU 设计团队'
    doc.core_properties.version = '1.1'
    doc.core_properties.comments = 'Design-review document and project scaffold publication; no MALU RTL implementation.'
    cover(doc)
    headings = []
    for line in lines:
        if line.startswith('## '):
            headings.append(line[3:].strip())
    if headings:
        doc.add_heading('阅读导航', level=1)
        for i, title in enumerate(headings, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            hyperlink(p, title, anchor='section_' + str(i))
        doc.add_page_break()
    i, ident, images, tables = 0, 0, [], 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), PALE)
            p._p.get_or_add_pPr().append(shd)
            rfont(p.add_run('\n'.join(block)), size=8.1, mono=True, color=NAVY)
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', v.replace(' ', '')) for v in row):
                    rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
                tables += 1
            continue
        if (m := re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)):
            image_path = source.parent / m.group(2)
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            add_figure(doc, image_path, m.group(1))
            images.append(image_path.name)
            i += 1
            continue
        if (m := re.match(r'^(#{1,4})\s+(.+)$', line)):
            depth, title = len(m.group(1)), m.group(2)
            if depth == 1 and 'MALU' in title:
                i += 1
                continue
            level = min(3, max(1, depth - 1))
            p = doc.add_heading(title, level=level)
            if depth == 2:
                ident += 1
                bookmark(p, ident)
            i += 1
            continue
        if line in ('---', '***'):
            i += 1
            continue
        if line.startswith('$$') or line.startswith(r'\['):
            formula = [line.replace('$$', '').replace(r'\[', '')]
            i += 1
            while i < len(lines) and not (lines[i].strip().endswith('$$') or lines[i].strip().endswith(r'\]')):
                formula.append(lines[i].strip())
                i += 1
            if i < len(lines):
                formula.append(lines[i].strip().replace('$$', '').replace(r'\]', ''))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, clean_formula(' '.join(formula)), size=10.3, base_color=NAVY)
            i += 1
            continue
        style = None
        if line.startswith('- '):
            style, line = 'List Bullet', line[2:]
        elif re.match(r'^\d+\.\s', line):
            style, line = 'List Number', re.sub(r'^\d+\.\s+', '', line)
        elif line.startswith('> '):
            line = line[2:]
        p = doc.add_paragraph(style=style)
        inline(p, line)
        i += 1
    figure_fallbacks = [
        ('01_malu_dot_datapath.png', '图：MALU 点积数据通路'),
        ('02_block_scale_groups.png', '图：BLOCK16 / BLOCK32 缩放映射'),
        ('03_transaction_and_backpressure.png', '图：事务与反压协议'),
        ('04_throughput_and_bandwidth.png', '图：吞吐与带宽约束'),
    ]
    missing = [(n, c) for n, c in figure_fallbacks if n not in images]
    if missing:
        doc.add_heading('工程图补充', level=1)
        for name, caption in missing:
            path = ROOT / 'figures' / name
            add_figure(doc, path, caption)
            images.append(name)
    doc.save(destination)
    with ZipFile(destination) as archive:
        assert archive.testzip() is None
        xml = archive.read('word/document.xml').decode('utf-8')
        assert 's1z4m11' in xml
        assert 'BLOCK16' in xml and 'BLOCK32' in xml
        assert len([n for n in archive.namelist() if n.startswith('word/media/')]) >= 4
    print({'file': str(destination), 'bytes': destination.stat().st_size,
           'headings': ident, 'tables': tables, 'figures': images})


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--source', type=Path, default=ROOT / 'MALU_design.md')
    ap.add_argument('--output', type=Path, default=ROOT / 'MALU_design_v1.1.docx')
    args = ap.parse_args()
    build(args.source, args.output)

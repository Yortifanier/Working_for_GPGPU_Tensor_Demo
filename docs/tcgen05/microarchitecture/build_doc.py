#!/usr/bin/env python3
"""Build the proposed tcgen05 Tensor microarchitecture DOCX/PDF from Markdown.

Pandoc retains native tables, links and Office Math. This file only formats
those objects; it never flattens equations to text or appends missing figures.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[2]
BUILD = REPO / "build/tcgen05_microarchitecture"
FONT = "Noto Sans CJK SC"
MONO = "DejaVu Sans Mono"
NAVY, TEAL, GRAY = "173553", "087F8C", "52687B"
NAME_PREFIX = "TCGEN05_Tensor_Microarchitecture_v"
FIGURES = ["01_tensor_subsystem", "02_dispatch_and_issue",
           "03_operand_and_tmem", "04_async_tracking", "05_malu_adapter",
           "06_smem_subsystem"]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}


def walk(value):
    if isinstance(value, dict):
        if "t" in value:
            yield value
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


def plain(value):
    if isinstance(value, list):
        return "".join(plain(v) for v in value)
    if not isinstance(value, dict):
        return ""
    kind, content = value.get("t"), value.get("c")
    if kind == "Str":
        return content
    if kind in {"Space", "SoftBreak", "LineBreak"}:
        return " "
    if kind in {"Code", "Math"}:
        return content[-1]
    if kind in {"Link", "Image"}:
        return plain(content[1])
    return plain(content)


def element(parent, tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    parent.append(node)
    return node


def property_element(parent,tag,**attrs):
    """Set a singleton formatting property without duplicating native OOXML."""
    nodes=parent.findall(qn(tag))
    node=nodes[0] if nodes else OxmlElement(tag)
    if not nodes:
        parent.append(node)
    for duplicate in nodes[1:]:
        parent.remove(duplicate)
    for key,value in attrs.items():
        node.set(qn("w:"+key),str(value))
    return node


def font(run, size=None, bold=None, color=None, mono=False):
    name = MONO if mono else FONT
    run.font.name = name
    props = run._r.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + key), FONT if key == "eastAsia" else name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_runs(paragraph):
    # Includes hyperlink runs, which Paragraph.runs alone does not expose.
    return [Run(node, paragraph) for node in paragraph._p.xpath(".//w:r")]


def get_style(doc,name):
    # Pandoc's display names may be capitalized differently from python-docx's
    # built-in aliases; use the actual style objects rather than alias lookup.
    return next((s for s in doc.styles if s.name.casefold()==name.casefold()),None)


def set_style(doc, name, size, color=None, bold=False):
    style=get_style(doc,name)
    if style is None:
        return
    style.font.name, style.font.size, style.font.bold = FONT, Pt(size), bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    props = style._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + key), FONT)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    element(run._r, "w:fldChar", fldCharType="begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    run._r.append(code)
    element(run._r, "w:fldChar", fldCharType="end")
    font(run, 8.5, color=GRAY)


def add_internal_link(paragraph, text, anchor):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = element(run, "w:rPr")
    element(props, "w:color", val=TEAL)
    element(props, "w:sz", val=21)
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)
    props.append(fonts)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def configure(doc, title, build_date, captions, version="1.1"):
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
        section.header_distance = section.footer_distance = Cm(.85)
        section.different_first_page_header_footer = True
    set_style(doc, "Normal", 10.5)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing, normal.space_after = 1.22, Pt(5)
    normal.widow_control = True
    for level, size in ((1,16), (2,12.5), (3,11.2), (4,10.5)):
        name = f"Heading {level}"
        set_style(doc, name, size, NAVY if level == 1 else TEAL, True)
        if get_style(doc,name) is not None:
            fmt = get_style(doc,name).paragraph_format
            fmt.keep_with_next, fmt.keep_together, fmt.widow_control = True, True, True
            fmt.space_before, fmt.space_after = Pt(13 if level == 1 else 9), Pt(6)
    if get_style(doc,"Caption") is None:
        doc.styles.add_style("Caption",WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Caption", "Image Caption", "Table Caption"):
        set_style(doc, name, 9.4, GRAY)
    set_style(doc, "Source Code", 9.0)
    set_style(doc, "Verbatim Char", 9.0)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        code = paragraph.style.name in {"Source Code", "Code"}
        if code:
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph_runs(paragraph):
            inline_code = run.style.name in {"Verbatim Char", "Code Char"}
            font(run, 9.0 if code or inline_code else None, mono=code or inline_code)
    if get_style(doc,"tcgen05 Table") is None:
        doc.styles.add_style("tcgen05 Table",WD_STYLE_TYPE.TABLE)
    for table in doc.tables:
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        table.style = "tcgen05 Table"
        props = table._tbl.tblPr
        borders = property_element(props,"w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            property_element(borders,"w:"+edge,val="single",sz=4,color="CBD7E1")
        widths = []
        for i in range(len(table.columns)):
            lengths = sorted(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1
                                 for ch in row.cells[i].text) for row in table.rows)
            typical = lengths[min(len(lengths)-1, int(len(lengths)*.65))]
            widths.append(max(7,min(60,typical))**.5)
        widths = [17*w/sum(widths) for w in widths]
        for column,width in zip(table.columns,widths):
            column.width = Cm(width)
        table_width = props.find(qn("w:tblW"))
        if table_width is not None:
            table_width.set(qn("w:type"), "dxa")
            table_width.set(qn("w:w"), str(round(Cm(17).twips)))
        for index,row in enumerate(table.rows):
            property_element(row._tr.get_or_add_trPr(), "w:cantSplit")
            if index == 0:
                property_element(row._tr.get_or_add_trPr(), "w:tblHeader")
            for i,cell in enumerate(row.cells):
                cell.width = Cm(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                cell_props = cell._tc.get_or_add_tcPr()
                property_element(cell_props,"w:shd",fill=NAVY if index == 0 else ("EDF5F8" if index%2 else "FFFFFF"))
                margins = property_element(cell_props,"w:tcMar")
                for edge in ("top","bottom","left","right"):
                    property_element(margins,"w:"+edge,w=65,type="dxa")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.12
                    paragraph.paragraph_format.keep_with_next = False
                    for run in paragraph_runs(paragraph):
                        font(run,9.0,bold=index == 0,color="FFFFFF" if index == 0 else "213342")
    image_paragraphs = []
    for shape in doc.inline_shapes:
        width,height = shape.width,shape.height
        # Pandoc may choose a smaller width; use all available width for the
        # architecture figures while retaining their original aspect.
        scale = min(Cm(17)/width,Cm(19)/height)
        shape.width,shape.height = int(width*scale),int(height*scale)
        node = shape._inline
        while node.tag != qn("w:p"):
            node = node.getparent()
        p = Paragraph(node,doc._body)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        image_paragraphs.append(p)
    for index,p in enumerate(image_paragraphs):
        next_node = p._p.getnext()
        next_p = Paragraph(next_node,doc._body) if next_node is not None and next_node.tag == qn("w:p") else None
        if next_p is None or "Caption" not in next_p.style.name:
            node = OxmlElement("w:p")
            p._p.addnext(node)
            next_p = Paragraph(node,doc._body)
            next_p.style = get_style(doc,"Caption")
            label = captions[index]
            next_p.add_run(label if re.match(r"图\s*\d",label) else f"图{index+1}  {label}")
        next_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        next_p.paragraph_format.keep_with_next = False
        next_p.paragraph_format.keep_together = True
        next_p.paragraph_format.space_after = Pt(9)
        for run in paragraph_runs(next_p):
            font(run,9.4,color=GRAY)
    # Native OMML equations keep their structure and a real installed math font.
    settings = doc.settings._element
    math_props = settings.find(qn("m:mathPr"))
    if math_props is None:
        math_props = OxmlElement("m:mathPr")
        settings.append(math_props)
    math_font = math_props.find(qn("m:mathFont"))
    if math_font is None:
        math_font = OxmlElement("m:mathFont")
        math_props.append(math_font)
    math_font.set(qn("m:val"), "DejaVu Math TeX Gyre")
    for section in doc.sections:
        header = section.header.paragraphs[0]
        font(header.add_run("tcgen05  /  PROJECT MICROARCHITECTURE"),8.5,color=GRAY)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        font(footer.add_run(f"v{version}  ·  "),8.5,color=GRAY)
        add_field(footer,"PAGE")
    doc.core_properties.title = title
    doc.core_properties.subject = "Proposed tcgen05 Tensor microarchitecture and existing RTL adaptation"
    doc.core_properties.author = "项目微架构研究记录"
    doc.core_properties.keywords = "tcgen05, proposed microarchitecture, TMEM, RTL adaptation, resource sharing"
    add_front_matter(doc,title,build_date,version,len(captions))


def add_front_matter(doc,title,build_date,version="1.1",figure_count=6):
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    identifiers = []
    existing = doc._element.xpath(".//w:bookmarkStart")
    next_id = max([int(n.get(qn("w:id"))) for n in existing]+[0])+1
    for i,p in enumerate(headings):
        name = "tcgen05_micro_section_"+str(i+1)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"),str(next_id+i));start.set(qn("w:name"),name)
        end = OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(next_id+i))
        p._p.insert(0,start);p._p.append(end)
        identifiers.append((p.text,name))
    prefix=[]
    def para(text="",size=10.5,bold=False,color=GRAY,space=8):
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(space)
        font(p.add_run(text),size,bold,color)
        prefix.append(p._p)
        return p
    para("PROJECT MICROARCHITECTURE",11,True,TEAL,22).paragraph_format.space_before=Pt(55)
    para(title,27,True,NAVY,18)
    para("拟议执行组织、资源复用与现有 RTL 适配",14,False,TEAL,28)
    para(f"版本 {version}  ·  "+build_date,11,False,GRAY,12)
    para(f"项目拟议方案；附 {figure_count} 张独立微架构示意图与可点击章节导航。",10.5,False,GRAY,8)
    para("现有 RTL 与待实现模块分别标注；本文件不是已实施 NVIDIA 微架构的描述。",10.5,False,GRAY,6)
    p=para();p.add_run().add_break(WD_BREAK.PAGE)
    para("阅读导航",18,True,NAVY,14).paragraph_format.keep_with_next=True
    for label,anchor in identifiers:
        # The 26 chapter links fit one page at the original 10.5 pt font.
        # Keep navigation spacing independent of the CJK body line metrics.
        p=para(space=3.5)
        p.paragraph_format.line_spacing=Pt(18)
        add_internal_link(p,label,anchor)
    p=para();p.add_run().add_break(WD_BREAK.PAGE)
    for index,node in enumerate(prefix):
        doc._body._body.remove(node)
        doc._body._body.insert(index,node)


def executable(name,explicit=None):
    if explicit:
        return str(Path(explicit).resolve())
    found=shutil.which(name)
    if found:
        return found
    if name == "soffice":
        candidate=Path("/opt/codex/runtimes/codex-primary-runtime/dependencies/bin/override/soffice")
        if candidate.is_file():
            return str(candidate)
    raise RuntimeError(f"Required program not found: {name}")


def normalized(text):
    return re.sub(r"\s+","",unicodedata.normalize("NFKC",text))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,default=ROOT/"TCGEN05_TENSOR_MICROARCHITECTURE.md")
    parser.add_argument("--date",default=datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat())
    parser.add_argument("--version",choices=("1.0","1.1"),default="1.1",
                        help="Output version; 1.1 requires six figures, 1.0 five")
    parser.add_argument("--title")
    parser.add_argument("--soffice")
    parser.add_argument("--docx-only",action="store_true")
    parser.add_argument("--render-check",action="store_true")
    args=parser.parse_args()
    name=NAME_PREFIX+args.version
    figures=FIGURES if args.version=="1.1" else FIGURES[:-1]
    build=BUILD/("v"+args.version)
    source=args.source.resolve()
    if not source.is_file():
        parser.error("Markdown source is not present yet: "+str(source))
    build.mkdir(parents=True,exist_ok=True)
    for directory in (build/"cache",build/"config"):
        directory.mkdir(exist_ok=True)
    env=dict(os.environ,XDG_CACHE_HOME=str(build/"cache"),XDG_CONFIG_HOME=str(build/"config"))
    source_bytes=source.read_bytes()
    (build/"document_qa.json").write_text(json.dumps({"passed":False,"status":"building", "document_version":args.version,"source_sha256":hashlib.sha256(source_bytes).hexdigest()},indent=2)+"\n")
    if "\ufffd" in source_bytes.decode("utf-8"):
        raise RuntimeError("Source contains Unicode replacement characters")
    pandoc=executable("pandoc")
    ast=json.loads(subprocess.check_output([pandoc,"-f","gfm+tex_math_dollars+implicit_figures","-t","json"],input=source_bytes,cwd=ROOT,env=env))
    title=args.title or "tcgen05 Tensor 微架构"
    if ast["blocks"] and ast["blocks"][0]["t"] == "Header" and ast["blocks"][0]["c"][0] == 1:
        first=ast["blocks"].pop(0)
        title=args.title or plain(first["c"][2])
    headings=[n for n in walk(ast["blocks"]) if n["t"] == "Header"]
    if headings:
        offset=min(n["c"][0] for n in headings)-1
        for n in headings:
            n["c"][0]-=offset
    images=[n for n in walk(ast["blocks"]) if n["t"] == "Image"]
    if len(images) != len(figures) or {Path(n["c"][2][0]).stem for n in images} != set(figures):
        raise RuntimeError(f"Version {args.version} source must explicitly reference each of the {len(figures)} microarchitecture figures once")
    for image in images:
        target=Path(image["c"][2][0])
        if target.suffix.lower() == ".svg" and (ROOT/target.with_suffix(".png")).is_file():
            image["c"][2][0]=str(target.with_suffix(".png"))
    figure_hashes={n["c"][2][0]:hashlib.sha256((ROOT/n["c"][2][0]).read_bytes()).hexdigest() for n in images}
    captions=[plain(n["c"][1]) for n in images]
    expected_math=sum(n["t"] == "Math" for n in walk(ast["blocks"]))
    expected_tables=sum(n["t"] == "Table" for n in walk(ast["blocks"]))
    ast_path=build/"microarchitecture_ast.json"
    ast_path.write_text(json.dumps(ast,ensure_ascii=False))
    body=build/"body.docx"
    subprocess.run([pandoc,str(ast_path),"-f","json","-t","docx","--resource-path",str(ROOT),"-o",str(body)],cwd=ROOT,env=env,check=True)
    doc=Document(body)
    configure(doc,title,args.date,captions,args.version)
    destination=ROOT/(name+".docx")
    staged=build/(name+".docx")
    doc.save(staged)
    with ZipFile(staged) as package:
        xml=ET.fromstring(package.read("word/document.xml"))
        media=[n for n in package.namelist() if n.startswith("word/media/")]
        actual_math=len(xml.findall(".//m:oMath",NS))
        actual_tables=len(xml.findall(".//w:tbl",NS))
        actual_navigation=sum((n.get(qn("w:anchor")) or "").startswith("tcgen05_micro_section_")
                              for n in xml.findall(".//w:hyperlink",NS))
        assert len(media)==len(figures),("media",media)
        assert {hashlib.sha256(package.read(name)).hexdigest() for name in media}==set(figure_hashes.values()),"embedded images differ from source PNGs"
        assert actual_math==expected_math,("OMML conversion",expected_math,actual_math)
        assert actual_tables==expected_tables,("native table conversion",expected_tables,actual_tables)
        assert actual_navigation==sum(n["c"][0]==1 for n in headings)
        assert "\ufffd" not in package.read("word/document.xml").decode()
    shutil.copy2(staged,destination)
    report={"title":title,"build_date":args.date,"document_version":args.version,"source":str(source.relative_to(REPO)),
            "source_sha256":hashlib.sha256(source_bytes).hexdigest(),"docx":str(destination.relative_to(REPO)),
            "figure_sha256":figure_hashes,
            "native_tables":actual_tables,"native_math":actual_math,"embedded_images":len(media),
            "chapter_navigation_links":actual_navigation,
            "pandoc_version":subprocess.check_output([pandoc,"--version"],text=True).splitlines()[0]}
    if not args.docx_only:
        export=build/("pdf_export_"+datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
        export.mkdir()
        profile=build/"libreoffice-profile"
        command=[executable("soffice",args.soffice),"-env:UserInstallation="+profile.as_uri(),"--headless","--convert-to","pdf:writer_pdf_Export","--outdir",str(export),str(destination)]
        result=subprocess.run(command,cwd=ROOT,env=env,text=True,capture_output=True,timeout=180)
        (build/"pdf_export.log").write_text(result.stdout+result.stderr)
        result.check_returncode()
        pdf=export/(name+".pdf")
        if not pdf.is_file() or pdf.stat().st_size<1000:
            raise RuntimeError(f"PDF was not generated; see {build.relative_to(REPO)}/pdf_export.log")
        pdf_dest=ROOT/pdf.name
        shutil.copy2(pdf,pdf_dest)
        info=subprocess.check_output([executable("pdfinfo"),str(pdf_dest)],text=True)
        pages=int(re.search(r"^Pages:\s+(\d+)",info,re.M)[1])
        text_path=build/"pdf_text.txt"
        subprocess.run([executable("pdftotext"),"-layout",str(pdf_dest),str(text_path)],check=True)
        text=text_path.read_text()
        if "\ufffd" in text:
            raise RuntimeError("PDF contains replacement characters")
        missing=[plain(n["c"][2]) for n in headings if normalized(plain(n["c"][2])) not in normalized(text)]
        if missing:
            raise RuntimeError("PDF lost heading text: "+str(missing))
        report.update(pdf=str(pdf_dest.relative_to(REPO)),pdf_pages=pages,pdf_text_characters=len(text),all_heading_text_found=True)
        (build/"pdfinfo.txt").write_text(info)
        try:
            import fitz
        except ImportError:
            report["pdf_structural_check"]="not run: optional PyMuPDF unavailable"
        else:
            with fitz.open(pdf_dest) as rendered:
                links=[link for page in rendered for link in page.get_links()]
                gotos=sum(link["kind"]==fitz.LINK_GOTO for link in links)
                image_count=len({image[0] for page in rendered for image in page.get_images()})
                outside=[]
                for number,page in enumerate(rendered,1):
                    for word in page.get_text("words"):
                        if word[0]<-1 or word[1]<-1 or word[2]>page.rect.width+1 or word[3]>page.rect.height+1:
                            outside.append([number,word[4]])
                if gotos<actual_navigation or image_count!=len(figures) or outside:
                    raise RuntimeError(f"PDF structure check failed: links={gotos}, images={image_count}, outside={outside}")
                report["pdf_structural_check"]={"internal_links":gotos,"unique_images":image_count,"text_outside_page_bounds":outside}
        if args.render_check:
            page_texts=text.split("\f")
            selected={1,pages}
            for caption in captions:
                token=normalized(caption)[:24]
                match=next((i+1 for i,page in enumerate(page_texts) if token and token in normalized(page)),None)
                if match:
                    selected.add(match)
            rendered=[]
            for number in sorted(selected):
                prefix=build/f"page_{number:02d}"
                subprocess.run([executable("pdftoppm"),"-f",str(number),"-l",str(number),"-singlefile","-scale-to","1600","-png",str(pdf_dest),str(prefix)],check=True)
                rendered.append(str(prefix.with_suffix(".png").relative_to(REPO)))
            report["rendered_pages"]=rendered
    report["source_unchanged_during_build"]=source.read_bytes()==source_bytes
    if not report["source_unchanged_during_build"]:
        raise RuntimeError("Source changed during build; rerun after editing is complete")
    if any(hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest for name,digest in figure_hashes.items()):
        raise RuntimeError("A source figure changed during build; rerun after editing is complete")
    report["docx_sha256"]=hashlib.sha256(destination.read_bytes()).hexdigest()
    if not args.docx_only:
        report["pdf_sha256"]=hashlib.sha256(pdf_dest.read_bytes()).hexdigest()
    report["passed"]=True
    report["status"]="complete"
    (build/"document_qa.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
